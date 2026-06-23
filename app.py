import os
import io
import csv
import re
import threading
import json
from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, Response
from flask_login import LoginManager, UserMixin, login_user, logout_user, login_required, current_user
from werkzeug.security import check_password_hash
from db import get_db, init_db
from scraper import scrape_jobs, load_all_plugins, get_favicon_url
from matcher import compute_match_score
import yaml

APP_VERSION = "0.5.0"
VERSION_HISTORY = [
    {"version": "0.5.0", "date": "2026-06-23", "changes": [
        "CV upload (PDF/DOCX/TXT) with keyword match scoring (0-10)",
        "Domain/tools/platforms/methods/languages weighted 2x, tech skills 1x",
        "Date range filter (last hour to 4 weeks + custom range)",
        "Remotive, Arbeitnow, Himalayas job portal plugins (69 total)",
        "Multi-select company filter with typeahead search",
        "Exclusion words filter",
        "Level and department filters",
    ]},
    {"version": "0.4.0", "date": "2026-06-22", "changes": [
        "Auto-scrape all sources on startup (cached for all users)",
        "Enhanced filters: job title, salary, location, work type, company",
        "Sort by company, posted date, location",
        "Job statuses: favorite, hidden, apply, applied, interview, rejected, withdrawn",
        "Filter by status tags",
        "Notes saved with each job",
        "CSV export of saved/tracked jobs",
    ]},
    {"version": "0.3.0", "date": "2026-06-22", "changes": [
        "60 plugins across 7 platforms for Munich tech companies",
        "Greenhouse, SmartRecruiters, Celonis API scrapers",
    ]},
    {"version": "0.2.0", "date": "2026-06-22", "changes": [
        "Workday + Oracle HCM API scrapers, pre-scrape filters, version display",
    ]},
    {"version": "0.1.0", "date": "2026-06-22", "changes": [
        "Initial release: Teamtailor/Personio scrapers, auth, plugin system",
    ]},
]

ALL_STATUSES = [
    "new", "viewed", "favorite", "hidden",
    "apply", "applied", "interview", "rejected", "withdrawn"
]

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "jobhunter-dev-key-change-in-prod")

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = "login"


@app.context_processor
def inject_globals():
    return {
        "app_version": APP_VERSION,
        "version_history": VERSION_HISTORY,
        "all_statuses": ALL_STATUSES,
    }


class User(UserMixin):
    def __init__(self, id, username):
        self.id = id
        self.username = username


@login_manager.user_loader
def load_user(user_id):
    db = get_db()
    row = db.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
    db.close()
    if row:
        return User(row["id"], row["username"])
    return None


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        db = get_db()
        row = db.execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
        db.close()
        if row and check_password_hash(row["password_hash"], password):
            login_user(User(row["id"], row["username"]))
            return redirect(url_for("feed"))
        flash("Invalid username or password.", "error")
    return render_template("login.html")


@app.route("/logout")
@login_required
def logout():
    logout_user()
    return redirect(url_for("login"))


@app.route("/")
@login_required
def feed():
    db = get_db()
    plugins = db.execute("SELECT * FROM plugins WHERE active = 1 ORDER BY name").fetchall()

    title_q = request.args.get("title", "").strip()
    location = request.args.get("location", "").strip()
    work_mode = request.args.get("work_mode", "")
    company = request.args.get("company", "")
    level = request.args.get("level", "")
    unit = request.args.get("unit", "").strip()
    min_salary = request.args.get("min_salary", type=int, default=0)
    status_filter = request.args.get("status", "")
    exclude_raw = request.args.get("exclude", "").strip()
    min_match = request.args.get("min_match", type=int, default=0)
    date_range = request.args.get("date_range", "")
    date_from = request.args.get("date_from", "")
    date_to = request.args.get("date_to", "")
    sort_by = request.args.get("sort", "date_desc")

    cv_row = db.execute("SELECT cv_text FROM user_cv WHERE user_id = ?", (current_user.id,)).fetchone()
    cv_text = cv_row["cv_text"] if cv_row else ""

    query = """
        SELECT j.*, p.name as source_name, p.base_url as source_url,
               COALESCE(uj.status, 'new') as user_status,
               COALESCE(uj.match_score, 0) as match_score,
               uj.notes, uj.applied_date, uj.follow_up_date
        FROM jobs j
        JOIN plugins p ON j.plugin_id = p.id
        LEFT JOIN user_jobs uj ON uj.job_id = j.id AND uj.user_id = ?
        WHERE 1=1
    """
    params = [current_user.id]

    if title_q:
        query += " AND LOWER(j.title) LIKE ?"
        params.append(f"%{title_q.lower()}%")
    if location:
        query += " AND LOWER(j.location) LIKE ?"
        params.append(f"%{location.lower()}%")
    if work_mode:
        query += " AND LOWER(j.work_mode) LIKE ?"
        params.append(f"%{work_mode.lower()}%")
    if min_salary:
        query += " AND j.salary_min >= ?"
        params.append(min_salary)
    if company:
        company_ids = [c.strip() for c in company.split(",") if c.strip()]
        if len(company_ids) == 1:
            query += " AND p.id = ?"
            params.append(int(company_ids[0]))
        elif company_ids:
            ph = ",".join("?" * len(company_ids))
            query += f" AND p.id IN ({ph})"
            params.extend(int(c) for c in company_ids)
    if level:
        query += " AND LOWER(j.seniority) LIKE ?"
        params.append(f"%{level.lower()}%")
    if unit:
        query += " AND LOWER(j.department) LIKE ?"
        params.append(f"%{unit.lower()}%")
    if exclude_raw:
        matches = re.findall(r'"([^"]+)"|([^,]+)', exclude_raw)
        exclude_terms = [(m[0] or m[1]).strip().lower() for m in matches]
        exclude_terms = [t for t in exclude_terms if t]
        for term in exclude_terms:
            query += " AND LOWER(j.title || ' ' || j.description || ' ' || j.department) NOT LIKE ?"
            params.append(f"%{term}%")
    if status_filter:
        if status_filter == "new":
            query += " AND (uj.status IS NULL OR uj.status = 'new')"
        elif status_filter == "hidden":
            query += " AND uj.status = 'hidden'"
        else:
            query += " AND uj.status = ?"
            params.append(status_filter)
    else:
        query += " AND (uj.status IS NULL OR uj.status != 'hidden')"

    date_map = {
        "1h": "-1 hour", "24h": "-1 day", "3d": "-3 days",
        "1w": "-7 days", "4w": "-28 days",
    }
    if date_range and date_range in date_map:
        query += f" AND j.first_seen >= datetime('now', ?)"
        params.append(date_map[date_range])
    elif date_range == "custom":
        if date_from:
            query += " AND j.first_seen >= ?"
            params.append(date_from)
        if date_to:
            query += " AND j.first_seen <= ?"
            params.append(date_to + " 23:59:59")

    sort_map = {
        "date_desc": "j.first_seen DESC",
        "date_asc": "j.first_seen ASC",
        "company_asc": "p.name ASC, j.first_seen DESC",
        "company_desc": "p.name DESC, j.first_seen DESC",
        "location_asc": "j.location ASC, j.first_seen DESC",
        "location_desc": "j.location DESC, j.first_seen DESC",
    }
    query += f" ORDER BY {sort_map.get(sort_by, 'j.first_seen DESC')}"

    raw_jobs = db.execute(query, params).fetchall()

    jobs = []
    for row in raw_jobs:
        j = dict(row)
        if cv_text and j.get("description"):
            job_text = f"{j['title']} {j['description']} {j.get('department', '')}"
            j["match_score"] = j["match_score"] or compute_match_score(job_text, cv_text)
        jobs.append(j)

    if min_match:
        jobs = [j for j in jobs if j.get("match_score", 0) >= min_match]

    if sort_by == "match_desc":
        jobs.sort(key=lambda j: j.get("match_score", 0), reverse=True)

    db.close()

    has_cv = bool(cv_text)
    filters = {
        "title": title_q, "location": location, "work_mode": work_mode,
        "company": company, "level": level, "unit": unit,
        "min_salary": min_salary or "", "status": status_filter,
        "exclude": exclude_raw, "min_match": min_match or "",
        "date_range": date_range, "date_from": date_from, "date_to": date_to,
        "sort": sort_by,
    }
    return render_template("feed.html", jobs=jobs, plugins=plugins,
                           get_favicon_url=get_favicon_url, filters=filters, has_cv=has_cv)


@app.route("/job/<int:job_id>/status", methods=["POST"])
@login_required
def update_job_status(job_id):
    data = request.get_json()
    status = data.get("status", "new")
    notes = data.get("notes")
    applied_date = data.get("applied_date", "")
    follow_up_date = data.get("follow_up_date", "")

    db = get_db()
    existing = db.execute(
        "SELECT notes FROM user_jobs WHERE user_id = ? AND job_id = ?",
        (current_user.id, job_id)).fetchone()
    if notes is None and existing:
        notes = existing["notes"]
    elif notes is None:
        notes = ""

    match_score = 0
    cv_row = db.execute("SELECT cv_text FROM user_cv WHERE user_id = ?", (current_user.id,)).fetchone()
    if cv_row and cv_row["cv_text"]:
        job_row = db.execute("SELECT title, description, department FROM jobs WHERE id = ?", (job_id,)).fetchone()
        if job_row:
            job_text = f"{job_row['title']} {job_row['description']} {job_row['department']}"
            match_score = compute_match_score(job_text, cv_row["cv_text"])

    db.execute("""
        INSERT INTO user_jobs (user_id, job_id, status, match_score, notes, applied_date, follow_up_date, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
        ON CONFLICT(user_id, job_id) DO UPDATE SET
            status=excluded.status, match_score=excluded.match_score, notes=excluded.notes,
            applied_date=excluded.applied_date, follow_up_date=excluded.follow_up_date,
            updated_at=CURRENT_TIMESTAMP
    """, (current_user.id, job_id, status, match_score, notes, applied_date, follow_up_date))
    db.commit()
    db.close()
    return jsonify({"ok": True})


@app.route("/scrape", methods=["POST"])
@login_required
def scrape():
    plugin_ids = request.form.getlist("plugin_ids")
    scrape_filters = {
        "keyword": request.form.get("scrape_keyword", "").strip(),
        "location": request.form.get("scrape_location", "").strip(),
        "work_mode": request.form.get("scrape_work_mode", "").strip(),
    }
    total_new, errors = run_scrape(plugin_ids, scrape_filters)
    if errors:
        flash(f"Scraped {total_new} new jobs. Errors: {'; '.join(errors[:5])}", "warning")
    else:
        flash(f"Scraped {total_new} new jobs.", "success")
    return redirect(url_for("feed"))


@app.route("/scrape/stream")
@login_required
def scrape_stream():
    plugin_ids = request.args.get("plugin_ids", "all")
    keyword = request.args.get("keyword", "")
    loc = request.args.get("location", "")
    wm = request.args.get("work_mode", "")
    user_id = current_user.id

    def generate():
        db = get_db()
        if plugin_ids == "all":
            rows = db.execute("SELECT * FROM plugins WHERE active = 1 ORDER BY name").fetchall()
        else:
            ids = [i.strip() for i in plugin_ids.split(",") if i.strip()]
            placeholders = ",".join("?" * len(ids))
            rows = db.execute(f"SELECT * FROM plugins WHERE id IN ({placeholders}) AND active = 1", ids).fetchall()

        total_plugins = len(rows)
        total_new = 0
        scrape_filters = {"keyword": keyword, "location": loc, "work_mode": wm}

        for idx, row in enumerate(rows):
            plugin_name = row["name"]
            yield f"data: {json.dumps({'type': 'progress', 'plugin': plugin_name, 'index': idx + 1, 'total_plugins': total_plugins, 'total_jobs': total_new})}\n\n"

            try:
                config = yaml.safe_load(row["config_yaml"])
                config["base_url"] = row["base_url"]
                jobs = scrape_jobs(config, scrape_filters=scrape_filters)
                plugin_new = 0
                new_job_ids = []
                for job in jobs:
                    try:
                        db.execute("""
                            INSERT INTO jobs (plugin_id, external_id, title, url, location, department,
                                             work_mode, employment_type, seniority, salary_text, description)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """, (row["id"], job["external_id"], job["title"], job["url"],
                              job.get("location", ""), job.get("department", ""), job.get("work_mode", ""),
                              job.get("employment_type", ""), job.get("seniority", ""),
                              job.get("salary_text", ""), job.get("description", "")))
                        plugin_new += 1
                        new_job_ids.append(db.execute("SELECT last_insert_rowid()").fetchone()[0])
                    except Exception:
                        pass
                db.commit()
                total_new += plugin_new

                if new_job_ids:
                    ph = ",".join("?" * len(new_job_ids))
                    new_jobs = db.execute(f"""
                        SELECT j.*, p.name as source_name, p.base_url as source_url,
                               'new' as user_status
                        FROM jobs j JOIN plugins p ON j.plugin_id = p.id
                        WHERE j.id IN ({ph})
                        ORDER BY j.id DESC
                    """, new_job_ids).fetchall()
                    jobs_data = [dict(j) for j in new_jobs]
                    yield f"data: {json.dumps({'type': 'jobs', 'plugin': plugin_name, 'new_count': plugin_new, 'total_jobs': total_new, 'jobs': jobs_data})}\n\n"
                else:
                    yield f"data: {json.dumps({'type': 'plugin_done', 'plugin': plugin_name, 'new_count': 0, 'total_jobs': total_new})}\n\n"

            except Exception as e:
                yield f"data: {json.dumps({'type': 'error', 'plugin': plugin_name, 'error': str(e)[:80], 'total_jobs': total_new})}\n\n"

        db.close()
        yield f"data: {json.dumps({'type': 'done', 'total_jobs': total_new})}\n\n"

    return Response(generate(), mimetype="text/event-stream",
                    headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


def run_scrape(plugin_ids=None, scrape_filters=None):
    db = get_db()
    if not plugin_ids or "all" in plugin_ids:
        rows = db.execute("SELECT * FROM plugins WHERE active = 1").fetchall()
    else:
        placeholders = ",".join("?" * len(plugin_ids))
        rows = db.execute(f"SELECT * FROM plugins WHERE id IN ({placeholders}) AND active = 1", plugin_ids).fetchall()

    total_new = 0
    errors = []
    for row in rows:
        try:
            config = yaml.safe_load(row["config_yaml"])
            config["base_url"] = row["base_url"]
            jobs = scrape_jobs(config, scrape_filters=scrape_filters)
            for job in jobs:
                try:
                    db.execute("""
                        INSERT INTO jobs (plugin_id, external_id, title, url, location, department,
                                         work_mode, employment_type, seniority, salary_text, description)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """, (row["id"], job["external_id"], job["title"], job["url"],
                          job.get("location", ""), job.get("department", ""), job.get("work_mode", ""),
                          job.get("employment_type", ""), job.get("seniority", ""),
                          job.get("salary_text", ""), job.get("description", "")))
                    total_new += 1
                except Exception:
                    pass
        except Exception as e:
            errors.append(f"{row['name']}: {str(e)[:80]}")
    db.commit()
    db.close()
    return total_new, errors


@app.route("/cv", methods=["GET", "POST"])
@login_required
def cv_page():
    db = get_db()
    if request.method == "POST":
        cv_text = ""
        if "cv_file" in request.files and request.files["cv_file"].filename:
            f = request.files["cv_file"]
            fname = f.filename.lower()
            raw = f.read()
            if fname.endswith(".pdf"):
                import fitz
                doc = fitz.open(stream=raw, filetype="pdf")
                cv_text = "\n".join(page.get_text() for page in doc)
                doc.close()
            elif fname.endswith((".docx", ".doc")):
                from docx import Document
                doc = Document(io.BytesIO(raw))
                cv_text = "\n".join(p.text for p in doc.paragraphs)
            elif fname.endswith(".txt"):
                cv_text = raw.decode("utf-8", errors="ignore")
            else:
                flash("Unsupported file format. Use PDF, DOCX, or TXT.", "error")
                return redirect(url_for("cv_page"))
            cv_filename = f.filename
        else:
            cv_text = request.form.get("cv_text", "").strip()
            cv_filename = "manual entry"

        if cv_text:
            db.execute("""
                INSERT INTO user_cv (user_id, cv_text, cv_filename, updated_at)
                VALUES (?, ?, ?, CURRENT_TIMESTAMP)
                ON CONFLICT(user_id) DO UPDATE SET
                    cv_text=excluded.cv_text, cv_filename=excluded.cv_filename,
                    updated_at=CURRENT_TIMESTAMP
            """, (current_user.id, cv_text, cv_filename))
            db.commit()
            flash(f"CV saved ({len(cv_text)} characters extracted).", "success")
        else:
            flash("No text could be extracted.", "error")

    cv = db.execute("SELECT * FROM user_cv WHERE user_id = ?", (current_user.id,)).fetchone()
    db.close()
    return render_template("cv.html", cv=cv)


@app.route("/tracker")
@login_required
def tracker():
    db = get_db()
    tracked_statuses = ("favorite", "apply", "applied", "interview", "rejected", "withdrawn")
    placeholders = ",".join("?" * len(tracked_statuses))
    jobs = db.execute(f"""
        SELECT j.*, p.name as source_name, p.base_url as source_url,
               uj.status, uj.notes, uj.applied_date, uj.follow_up_date, uj.updated_at as status_updated
        FROM user_jobs uj
        JOIN jobs j ON uj.job_id = j.id
        JOIN plugins p ON j.plugin_id = p.id
        WHERE uj.user_id = ? AND uj.status IN ({placeholders})
        ORDER BY uj.updated_at DESC
    """, (current_user.id, *tracked_statuses)).fetchall()
    db.close()
    return render_template("tracker.html", jobs=jobs, get_favicon_url=get_favicon_url)


@app.route("/tracker/export")
@login_required
def export_tracker():
    db = get_db()
    tracked_statuses = ("favorite", "apply", "applied", "interview", "rejected", "withdrawn")
    placeholders = ",".join("?" * len(tracked_statuses))
    jobs = db.execute(f"""
        SELECT j.title, j.url, j.location, j.work_mode, j.employment_type, j.department,
               j.salary_text, j.first_seen, p.name as company,
               uj.status, uj.notes, uj.applied_date, uj.follow_up_date, uj.updated_at
        FROM user_jobs uj
        JOIN jobs j ON uj.job_id = j.id
        JOIN plugins p ON j.plugin_id = p.id
        WHERE uj.user_id = ? AND uj.status IN ({placeholders})
        ORDER BY uj.updated_at DESC
    """, (current_user.id, *tracked_statuses)).fetchall()
    db.close()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Title", "Company", "URL", "Location", "Work Mode", "Employment Type",
                     "Department", "Salary", "Status", "Notes", "Applied Date",
                     "Follow-up Date", "First Seen", "Last Updated"])
    for j in jobs:
        writer.writerow([j["title"], j["company"], j["url"], j["location"], j["work_mode"],
                         j["employment_type"], j["department"], j["salary_text"],
                         j["status"], j["notes"], j["applied_date"],
                         j["follow_up_date"], j["first_seen"], j["updated_at"]])

    return Response(
        output.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=jobhunter_export.csv"})


@app.route("/plugins")
@login_required
def plugins_page():
    db = get_db()
    plugins = db.execute("SELECT * FROM plugins ORDER BY name").fetchall()
    db.close()
    return render_template("plugins.html", plugins=plugins)


@app.route("/plugins/upload", methods=["POST"])
@login_required
def upload_plugin():
    if "file" not in request.files:
        flash("No file selected.", "error")
        return redirect(url_for("plugins_page"))

    f = request.files["file"]
    if not f.filename.endswith((".yaml", ".yml")):
        flash("Only YAML files are accepted.", "error")
        return redirect(url_for("plugins_page"))

    content = f.read().decode("utf-8")
    try:
        config = yaml.safe_load(content)
    except Exception:
        flash("Invalid YAML.", "error")
        return redirect(url_for("plugins_page"))

    if not config.get("name") or not config.get("base_url"):
        flash("Plugin must have 'name' and 'base_url' fields.", "error")
        return redirect(url_for("plugins_page"))

    from werkzeug.utils import secure_filename
    filename = secure_filename(f.filename)

    plugin_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "plugins")
    filepath = os.path.join(plugin_dir, filename)
    with open(filepath, "w", encoding="utf-8") as pf:
        pf.write(content)

    db = get_db()
    db.execute("""
        INSERT INTO plugins (filename, name, platform, base_url, config_yaml, active)
        VALUES (?, ?, ?, ?, ?, 1)
        ON CONFLICT(filename) DO UPDATE SET
            name=excluded.name, platform=excluded.platform,
            base_url=excluded.base_url, config_yaml=excluded.config_yaml
    """, (filename, config["name"], config.get("platform", "generic"),
          config["base_url"], content))
    db.commit()
    db.close()

    flash(f"Plugin '{config['name']}' uploaded.", "success")
    return redirect(url_for("plugins_page"))


@app.route("/plugins/<int:plugin_id>/toggle", methods=["POST"])
@login_required
def toggle_plugin(plugin_id):
    db = get_db()
    db.execute("UPDATE plugins SET active = CASE WHEN active=1 THEN 0 ELSE 1 END WHERE id = ?", (plugin_id,))
    db.commit()
    db.close()
    return redirect(url_for("plugins_page"))


def seed_db():
    from werkzeug.security import generate_password_hash
    db = get_db()
    users = [("haroon", "haroon123"), ("broon", "broon123"), ("croon", "croon123")]
    for username, password in users:
        existing = db.execute("SELECT id FROM users WHERE username = ?", (username,)).fetchone()
        if not existing:
            db.execute("INSERT INTO users (username, password_hash) VALUES (?, ?)",
                       (username, generate_password_hash(password)))
    plugins_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "plugins")
    for fname in os.listdir(plugins_dir):
        if not fname.endswith((".yaml", ".yml")):
            continue
        with open(os.path.join(plugins_dir, fname), "r", encoding="utf-8") as f:
            content = f.read()
            config = yaml.safe_load(content)
        existing = db.execute("SELECT id FROM plugins WHERE filename = ?", (fname,)).fetchone()
        if not existing:
            db.execute(
                "INSERT INTO plugins (filename, name, platform, base_url, config_yaml, active) VALUES (?,?,?,?,?,1)",
                (fname, config["name"], config.get("platform", "generic"), config["base_url"], content))
    db.commit()
    db.close()


def startup_scrape():
    """Background scrape of all active plugins on app start."""
    with app.app_context():
        db = get_db()
        job_count = db.execute("SELECT COUNT(*) as c FROM jobs").fetchone()["c"]
        db.close()
        if job_count > 0:
            print(f"[startup] {job_count} cached jobs found, skipping initial scrape.")
            return
        print("[startup] No cached jobs, scraping all active plugins...")
        total, errors = run_scrape(["all"])
        print(f"[startup] Scraped {total} new jobs. Errors: {len(errors)}")


with app.app_context():
    init_db()
    seed_db()
    threading.Thread(target=startup_scrape, daemon=True).start()


if __name__ == "__main__":
    app.run(debug=True, port=5000)
